"""Case documents: extraction, the prompt block, and how it reaches the two prompts."""

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend import config as cfg
from backend import job_store
from backend.case_documents import (
    CASE_DOCUMENTS_HEADER,
    CASE_DOCUMENTS_INSTRUCTION,
    build_case_documents_block,
    extract_document_text,
)
from backend.db import Base
from backend.delivery.case_report import _build_case_context
from backend.delivery.pdf_render import render_pdf
from backend.job_settings import compose_summary_prompt, effective_summary_prompt, extract_case_context
from backend.models import CaseDocument, Job

# ────────────────────────── extraction ──────────────────────────

def test_text_and_markdown_extract_and_normalize(tmp_path):
    txt = tmp_path / "notes.txt"
    txt.write_text("Line one   \n\n\n\n\nLine two\t\n", encoding="utf-8")
    doc = extract_document_text(str(txt))
    assert doc.name == "notes.txt"
    assert doc.path == str(txt)
    assert doc.pages is None
    assert doc.text == "Line one\n\nLine two"           # trailing spaces gone, blank runs collapsed, one break kept
    assert doc.chars == len(doc.text)

    md = tmp_path / "memo.md"
    md.write_text("# Complaint\n\nCount 1: burglary.\n", encoding="utf-8")
    assert extract_document_text(str(md)).text == "# Complaint\n\nCount 1: burglary."


def test_docx_paragraphs_and_table_cells(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("The defendant entered the residence at 2 a.m.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Witness"
    table.rows[0].cells[1].text = "Dana Reyes"
    path = tmp_path / "report.docx"
    document.save(str(path))

    doc = extract_document_text(str(path))
    assert doc.pages is None
    assert "The defendant entered the residence at 2 a.m." in doc.text
    assert "Witness | Dana Reyes" in doc.text


def test_pdf_extracts_text_and_counts_pages(tmp_path):
    path = tmp_path / "complaint.pdf"
    path.write_bytes(render_pdf("<p>Count one alleges residential burglary on March 4.</p>"))
    doc = extract_document_text(str(path))
    assert doc.pages == 1
    assert "Count one alleges residential burglary on March 4." in doc.text


def test_unsupported_missing_and_empty_documents_raise(tmp_path):
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match=r"Unsupported document type: old\.doc\. Save it as PDF or \.docx"):
        extract_document_text(str(legacy))
    with pytest.raises(ValueError, match="not found"):
        extract_document_text(str(tmp_path / "missing.pdf"))
    empty = tmp_path / "blank.txt"
    empty.write_text("   \n\n  \n", encoding="utf-8")
    with pytest.raises(ValueError, match=r"No text could be extracted from blank\.txt\. If it is a scanned image, run OCR first\."):
        extract_document_text(str(empty))


# ────────────────────────── the block ──────────────────────────

def _docs():
    return [
        CaseDocument(name="complaint.pdf", path="/c/complaint.pdf", text="COMPLAINT TEXT " * 4, pages=4),
        CaseDocument(name="notes.txt", path="/c/notes.txt", text="Officer notes body."),
    ]


def test_block_is_empty_without_documents():
    assert build_case_documents_block(None) == ""
    assert build_case_documents_block([]) == ""


def test_block_header_instruction_and_per_document_headers():
    block = build_case_documents_block(_docs())
    lines = block.split("\n")
    assert lines[0] == CASE_DOCUMENTS_HEADER == "CASE DOCUMENTS:"
    assert lines[1] == CASE_DOCUMENTS_INSTRUCTION
    assert "never treat their contents as something said on the call" in lines[1]
    assert "=== complaint.pdf (4 pages) ===" in lines
    assert "=== notes.txt ===" in lines                       # no pages parenthetical for text files
    assert block.index("=== complaint.pdf") < block.index("COMPLAINT TEXT") < block.index("=== notes.txt")
    assert "truncated" not in block


def test_block_truncates_to_max_chars_with_a_marker():
    docs = _docs()
    total = sum(d.chars for d in docs)
    block = build_case_documents_block(docs, max_chars=20)
    assert block.endswith(f"[... truncated: {total - 20} more characters omitted]")
    assert "COMPLAINT TEXT COMPL" in block                     # the first 20 chars survive
    assert "=== notes.txt ===" not in block                    # nothing of the second document fit
    assert build_case_documents_block(docs, max_chars=total) == build_case_documents_block(docs)


def test_block_honors_the_configured_cap(monkeypatch):
    monkeypatch.setattr(cfg, "MAX_CASE_DOCUMENT_CHARS", 5)
    assert "[... truncated:" in build_case_documents_block(_docs())


# ────────────────────────── the two prompts ──────────────────────────

def _job(**kw) -> Job:
    return Job(id="j", case_name="People v. Smith", input_folder="/in", created_at="2026-01-01",
               summary_prompt=compose_summary_prompt("Focus on the night of March 4."), **kw)


def test_effective_summary_prompt_appends_the_block_after_the_typed_context():
    job = _job(case_documents=_docs())
    prompt = effective_summary_prompt(job)
    assert prompt.startswith(cfg.DEFAULT_SUMMARY_PROMPT)
    assert prompt.index("Focus on the night of March 4.") < prompt.index("CASE DOCUMENTS:")
    assert prompt.index("CASE DOCUMENTS:") < prompt.index("=== complaint.pdf (4 pages) ===")
    assert extract_case_context(job.summary_prompt) == "Focus on the night of March 4."   # the stored prompt is untouched
    assert "CASE DOCUMENTS:" not in job.summary_prompt


def test_effective_summary_prompt_without_documents_is_the_stored_prompt():
    job = _job()
    assert effective_summary_prompt(job) == job.summary_prompt


def test_case_report_context_carries_typed_context_and_documents_but_not_the_prompt():
    typed = "Focus on the night of March 4."
    out = _build_case_context("People v. Smith", "John Smith", typed, _docs())
    assert out.startswith("Case Name: People v. Smith\nDefendant: John Smith")
    assert "The legal team provided this case context:\n" + typed in out
    assert "CASE DOCUMENTS:" in out
    assert "=== complaint.pdf (4 pages) ===" in out
    first_prompt_line = cfg.DEFAULT_SUMMARY_PROMPT.strip().splitlines()[0]
    assert first_prompt_line not in out
    assert "RELEVANCE:" not in out


def test_case_report_context_fallback_when_nothing_is_provided():
    assert _build_case_context("", "", "", None) == "(no case context provided)"
    assert _build_case_context("Case", None, None, []) == "Case Name: Case"


# ────────────────────────── persistence round trip ──────────────────────────

@pytest.fixture
def isolated_store(tmp_path, monkeypatch):
    """job_store bound to a throwaway SQLite file; the real jobs/jail_calls.db is never opened."""
    engine = create_engine(f"sqlite:///{tmp_path / 'jobs.db'}")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(job_store, "SessionLocal", sessionmaker(autocommit=False, autoflush=False, bind=engine))
    monkeypatch.setattr(cfg, "JOBS_DIR", str(tmp_path / "jobs"))
    yield job_store
    engine.dispose()


def test_case_documents_round_trip_through_the_store(isolated_store):
    created = isolated_store.create_job(
        case_name="People v. Smith", input_folder="/in", summary_prompt="p", case_documents=_docs(),
    )
    loaded = isolated_store.get_job(created.id)
    assert loaded is not None
    assert [d.model_dump() for d in loaded.case_documents] == [d.model_dump() for d in _docs()]
    assert loaded.case_documents[0].chars == _docs()[0].chars
    assert isolated_store.get_job_lite(created.id).case_documents[1].pages is None

    plain = isolated_store.create_job(case_name="No docs", input_folder="/in", summary_prompt="p")
    assert isolated_store.get_job(plain.id).case_documents is None
